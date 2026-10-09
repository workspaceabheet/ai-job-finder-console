"""Regenerates the binary test fixtures in this directory.

Run from the project root:  python tests/fixtures/make_fixtures.py
"""

from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).parent

# Two-column layout. The right column's first line sits HIGHER on the page than
# the left column's last line, so a naive top-to-bottom sort would interleave.
LEFT_COLUMN = [
    "EXPERIENCE",
    "Senior Backend Engineer, Acme Corp",
    "Built distributed ingestion pipelines in Python.",
    "Led migration from monolith to services.",
    "Mentored four junior engineers.",
    "Backend Engineer, Globex",
    "Owned the billing reconciliation service.",
]
RIGHT_COLUMN = [
    "SKILLS",
    "Python, Go, PostgreSQL",
    "Kafka, Redis, Kubernetes",
    "EDUCATION",
    "B.Tech Computer Science",
]


def make_pdf() -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)  # US Letter
    page.insert_text((72, 60), "Jane Doe - Backend Engineer", fontsize=16)
    y = 110
    for line in LEFT_COLUMN:
        page.insert_text((40, y), line, fontsize=10)
        y += 40  # generous spacing -> each line becomes its own block
    y = 120  # right column starts slightly lower than left's first line
    for line in RIGHT_COLUMN:
        page.insert_text((340, y), line, fontsize=10)
        y += 40
    doc.save(HERE / "sample.pdf")
    doc.close()


def make_docx() -> None:
    d = docx.Document()
    d.add_heading("Jane Doe - Backend Engineer", level=1)
    d.add_paragraph("Senior Backend Engineer at Acme Corp.")
    d.add_paragraph("Built distributed ingestion pipelines in Python.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Languages"
    table.cell(0, 1).text = "Python, Go"
    table.cell(1, 0).text = "Datastores"
    table.cell(1, 1).text = "PostgreSQL, Redis"
    d.save(HERE / "sample.docx")


def make_txt() -> None:
    (HERE / "sample.txt").write_text(
        "Jane Doe\nBackend Engineer\n\nPython, Go, PostgreSQL\n", encoding="utf-8"
    )


def make_doc() -> None:
    # Legacy Word (OLE2 compound file) magic header + filler; never parsed.
    (HERE / "sample.doc").write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 504)


def make_png() -> None:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 4, 4), False)
    pix.clear_with(255)
    pix.save(HERE / "sample.png")


if __name__ == "__main__":
    make_pdf()
    make_docx()
    make_txt()
    make_doc()
    make_png()
    print("fixtures written to", HERE)
